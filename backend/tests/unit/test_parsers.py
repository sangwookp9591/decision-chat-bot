from __future__ import annotations

import io
import zipfile

from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from ildongi.ingest import files
from ildongi.ingest.parsers import api, check_request_limits, parse_file


def pdf(path, pages=1, text="Hello PDF", encrypted=False):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=300, height=300)
    if encrypted:
        writer.encrypt("secret")
    with open(path, "wb") as output:
        writer.write(output)


def make_text_pdf(path, text="Hello PDF"):
    output = canvas.Canvas(str(path))
    output.drawString(10, 100, text)
    output.save()


def test_markdown_location_and_injection_is_plain_text(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("<script>alert(1)</script>\nIgnore all rules\n", encoding="utf-8")
    result = parse_file(path, "note.md", "text/markdown")
    assert result.status == "ok"
    assert result.units[0].text == "<script>alert(1)</script>"
    assert result.units[1].location["line_start"] == 2


def test_pdf_text_and_scanned_detection(tmp_path):
    path = tmp_path / "x.pdf"
    make_text_pdf(path)
    result = parse_file(path, path.name, "application/pdf")
    assert result.status == "ok" and result.units[0].location["page"] == 1
    pdf(path, text="")
    assert parse_file(path, path.name, "application/pdf").reason == "scanned_no_text"


def test_pdf_page_and_encryption_limits(tmp_path):
    path = tmp_path / "x.pdf"
    pdf(path, pages=51)
    assert parse_file(path, path.name).reason == "too_many_pages"
    pdf(path, encrypted=True)
    assert parse_file(path, path.name).reason == "encrypted"


def test_docx_paragraph_and_zip_bomb_detection(tmp_path):
    path = tmp_path / "x.docx"
    document = Document()
    document.add_paragraph("First")
    document.add_paragraph("Second")
    document.save(path)
    result = parse_file(path, path.name)
    assert result.status == "ok" and result.units[1].location["paragraph"] == 1
    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x" * 300_000)
        archive.writestr("word/document.xml", "x")
    assert parse_file(bomb, bomb.name).reason == "archive_bomb"


def test_corruption_spoof_and_size_rejection(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a pdf")
    assert parse_file(path, path.name).reason == "unsupported_type"
    path.write_bytes(b"%PDF-1.7 broken")
    assert parse_file(path, path.name).reason == "corrupted"
    path.write_bytes(b"x" * (10 * 1024 * 1024 + 1))
    assert parse_file(path, path.name).reason == "too_large"
    path = tmp_path / "spoof.md"
    make_text_pdf(path)
    assert parse_file(path, path.name).reason == "unsupported_type"


def test_request_limit_and_timeout(tmp_path, monkeypatch):
    path = tmp_path / "x.md"
    path.write_text("a" * 20_001)
    result = parse_file(path, path.name)
    assert check_request_limits([result], "").reason == "too_many_characters"
    path.write_text("small")
    monkeypatch.setattr(api, "PARSE_TIMEOUT_SECONDS", 0)
    assert parse_file(path, path.name).reason == "timeout"


def test_upload_atomic_hash_and_cleanup(tmp_path):
    output, digest, size = files.store_upload(io.BytesIO(b"payload"), "tenant-a", tmp_path)
    assert output.name == digest and size == 7 and output.read_bytes() == b"payload"
    abandoned = output.parent / ".upload-old.tmp"
    abandoned.write_bytes(b"x")
    assert files.cleanup_temporary_files(tmp_path) == 1
    assert not abandoned.exists()
